"""Import real contents before treating a source as readable evidence."""
from __future__ import annotations
from html.parser import HTMLParser
from io import BytesIO
import asyncio
import ipaddress
import re
import socket
from urllib.parse import urljoin, urlparse
from zipfile import ZipFile

import httpx
from app.essay.models import SourceRecord, identifier, now
from app.essay.citations import content_hash

MAX_BYTES = 12 * 1024 * 1024
MAX_TEXT = 250000


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skip = 0
        self.title = []
        self.in_title = False
        self.authors = []
        self.published = ""

    def handle_starttag(self, tag, attrs):
        if tag == "meta":
            values = dict(attrs)
            name = (values.get("name") or values.get("property") or "").lower()
            value = values.get("content") or ""
            if name in {"citation_author", "author", "dc.creator", "dc.contributor"} and value:
                if value not in self.authors:
                    self.authors.append(value)
            if name in {"citation_publication_date", "article:published_time", "dc.date", "dc.date.issued"} and not self.published:
                self.published = value
        if tag in {"script", "style", "noscript", "svg", "nav", "footer"}:
            self.skip += 1
        if tag == "title":
            self.in_title = True
        if tag in {"p", "div", "br", "h1", "h2", "h3", "li", "article", "section"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg", "nav", "footer"} and self.skip:
            self.skip -= 1
        if tag == "title":
            self.in_title = False
        if tag in {"p", "div", "li"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
        elif not self.skip:
            self.parts.append(data)

    def text(self):
        return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n+", "\n\n", "".join(self.parts))).strip()


def decode_document(data, filename, media_type=""):
    if len(data) > MAX_BYTES:
        raise ValueError("Il documento supera il limite di 12 MB.")
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "pdf" or "application/pdf" in media_type:
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Il PDF è protetto: importa una copia testuale accessibile.")
        text = "\n\n".join("[Pagina " + str(i + 1) + "]\n" + (p.extract_text() or "")
                            for i, p in enumerate(reader.pages))
        if sum(len(p.extract_text() or "") for p in reader.pages) < 40:
            raise ValueError("Il PDF non contiene testo leggibile; per questa scansione serve OCR.")
    elif suffix == "docx":
        from docx import Document
        with ZipFile(BytesIO(data)) as archive:
            if sum(e.file_size for e in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Il contenuto DOCX è troppo grande.")
        document = Document(BytesIO(data))
        text = "\n\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows)
    elif suffix in {"txt", "md", "markdown"} or media_type.startswith("text/plain"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("cp1252")
    elif "text/html" in media_type or suffix in {"html", "htm"}:
        parser = TextHTML()
        parser.feed(data.decode("utf-8", errors="replace"))
        text = parser.text()
    else:
        raise ValueError("Formato non supportato. Usa PDF testuale, DOCX, TXT o Markdown.")
    if len(text.strip()) < 20:
        raise ValueError("Il documento non contiene abbastanza testo leggibile.")
    if len(text) > MAX_TEXT:
        raise ValueError("Il testo supera 250.000 caratteri: importa una sezione del documento.")
    return text


async def validate_public_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in {"https", "http"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Inserisci un URL pubblico HTTP o HTTPS.")
    if parsed.port not in (None, 80, 443):
        raise ValueError("La fonte deve usare una porta web standard.")
    addresses = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, parsed.port or 443)
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise ValueError("La ricerca acquisisce solo fonti web pubbliche.")
    return url


async def acquire_url(url, title=""):
    await validate_public_url(url)
    async with httpx.AsyncClient(timeout=25, follow_redirects=False, trust_env=False,
                                 headers={"User-Agent": "SaggIA/0.1 (local research studio)"}) as client:
        for _ in range(6):
            async with client.stream("GET", url) as response:
                if response.is_redirect:
                    url = urljoin(url, response.headers.get("location", ""))
                    await validate_public_url(url)
                    continue
                response.raise_for_status()
                chunks, size = [], 0
                async for chunk in response.aiter_bytes():
                    size += len(chunk)
                    if size > MAX_BYTES:
                        raise ValueError("La fonte supera 12 MB.")
                    chunks.append(chunk)
                data = b"".join(chunks)
                media = response.headers.get("content-type", "")
                name = urlparse(url).path.rsplit("/", 1)[-1] or "pagina.html"
                text = await asyncio.to_thread(decode_document, data, name, media)
                author, published = "", ""
                if "html" in media:
                    parser = TextHTML()
                    parser.feed(data.decode("utf-8", errors="replace"))
                    title = "".join(parser.title).strip() or title
                    author, published = "; ".join(parser.authors), parser.published
                return {"url": url, "title": title or name, "text": text, "author": author, "published": published}
        raise ValueError("Troppi reindirizzamenti nella fonte.")


def source_from_text(title, text, **metadata):
    if len(text.strip()) < 20 or len(text) > MAX_TEXT:
        raise ValueError("Inserisci un testo tra 20 e 250.000 caratteri.")
    return SourceRecord(id=identifier("S"), title=title, text=text, status="acquired",
                        content_hash=content_hash(text), acquired_at=now(), **metadata).model_dump()
