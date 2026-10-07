"""Canonical citation rendering for PDF, EPUB and DOCX exports."""
from __future__ import annotations
from html import escape
from io import BytesIO
import re
import tempfile
from pathlib import Path
import markdown
from app.essay.citations import rendered_manuscript

def safe_html(text):
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    html = markdown.markdown(text.replace("<", "&lt;"), extensions=["tables"])
    html = re.sub(r'href="(?:javascript|data):[^"]*"', 'href="#"', html, flags=re.I)
    return html

def text_blocks(text):
    for block in re.split(r"\n\s*\n", text.strip()):
        block = block.strip()
        if block:
            yield block

def plain(text):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", re.sub(r"[*_\x60]", "", text))

def export_book(session, format):
    chapters, bibliography = rendered_manuscript(session)
    title = session.form_data.title or session.plan["title"]
    subtitle = session.form_data.subtitle or session.plan.get("subtitle", "")
    author = session.form_data.author_name
    work_note = "Versione di lavoro. Consulta il rapporto di revisione nello studio." if session.review is None or session.review.get("partial") or session.review.get("issues") else ""
    if format == "pdf":
        from app.essay.pdf_layout import export_pdf
        return export_pdf(session, chapters, bibliography, title, subtitle, author, work_note)
    if format == "docx":
        from docx import Document
        from docx.shared import Pt, Inches
        doc = Document()
        doc.core_properties.title = title[:255]
        doc.core_properties.author = author[:255]
        if session.cover_image_path and Path(session.cover_image_path).exists():
            doc.add_picture(session.cover_image_path, width=Inches(4.5))
            doc.add_page_break()
        doc.add_heading(title, 0)
        if subtitle:
            doc.add_paragraph(subtitle, "Subtitle")
        if author:
            doc.add_paragraph(author)
        doc.add_page_break()
        doc.add_heading("Indice", 1)
        for chapter in chapters:
            doc.add_paragraph(chapter["title"])
        if bibliography:
            doc.add_paragraph("Bibliografia")
        for chapter in chapters:
            doc.add_page_break()
            doc.add_heading(chapter["title"], 1)
            for block in text_blocks(chapter["display_content"]):
                if block.startswith("#"):
                    level = min(4, len(block) - len(block.lstrip("#")) + 1)
                    doc.add_heading(plain(block.lstrip("# ")), level)
                else:
                    doc.add_paragraph(plain(block))
        if bibliography:
            doc.add_page_break()
            doc.add_heading("Bibliografia", 1)
            for source in bibliography:
                doc.add_paragraph("[" + source["label"] + "] " + source["text"])
        if work_note:
            doc.add_paragraph(work_note)
        doc.styles["Normal"].font.name = "Georgia"
        doc.styles["Normal"].font.size = Pt(11)
        buffer = BytesIO()
        doc.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if format == "epub":
        from ebooklib import epub
        book = epub.EpubBook()
        book.set_identifier(session.session_id)
        book.set_title(title)
        book.set_language("it")
        if author:
            book.add_author(author)
        if session.cover_image_path and Path(session.cover_image_path).exists():
            book.set_cover("cover.png", Path(session.cover_image_path).read_bytes())
        front = epub.EpubHtml(title=title, file_name="title.xhtml", lang="it")
        front.content = "<h1>" + escape(title) + "</h1><p>" + escape(subtitle) + "</p><p>" + escape(author) + "</p>"
        book.add_item(front)
        sections = [front]
        for index, chapter in enumerate(chapters):
            item = epub.EpubHtml(title=chapter["title"], file_name="chapter-" + str(index) + ".xhtml", lang="it")
            item.content = "<h1>" + escape(chapter["title"]) + "</h1>" + safe_html(chapter["display_content"])
            book.add_item(item)
            sections.append(item)
        if bibliography:
            item = epub.EpubHtml(title="Bibliografia", file_name="bibliography.xhtml", lang="it")
            item.content = "<h1>Bibliografia</h1>" + "".join("<p>" + escape("[" + s["label"] + "] " + s["text"]) + "</p>" for s in bibliography)
            book.add_item(item)
            sections.append(item)
        if work_note:
            item = epub.EpubHtml(title="Nota editoriale", file_name="work-note.xhtml", lang="it")
            item.content = "<p>" + escape(work_note) + "</p>"
            book.add_item(item)
            sections.append(item)
        book.toc = sections
        book.spine = ["nav", *sections]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "essay.epub"
            epub.write_epub(str(target), book)
            return target.read_bytes(), "application/epub+zip"
    raise ValueError("Formato non supportato.")


def pdf_page_count(session):
    from pypdf import PdfReader
    data, _ = export_book(session, "pdf")
    return len(PdfReader(BytesIO(data)).pages)
